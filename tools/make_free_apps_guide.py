"""Free_Apps_Activity_Guide_Ch1-7.docx/.pdf + Import_Files zip, from make_free_apps.py output."""
import os, sys, subprocess, zipfile, glob
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from make_activity_pack import shade, borders, widths, field, para, label, heading, NAVY, TEAL
from activities import A
from free_apps import CH, MAP, WORDWALL, PADLET, MENTI
import make_free_apps as M
OUT = M.OUT; NAME = 'Free_Apps_Activity_Guide_Ch1-7'
REL = lambda p: os.path.relpath(p, OUT)
CHECKED = 'Checked on the official help and pricing pages on Oct 16, 2026. Plans change; check again before class.'

APPS = [
 dict(app='Kahoot', plan='Kahoot! Go (free)', limit='10 players per live game or assignment. Import works for standard quiz questions only, 4 answers, question 120 / answer 75 characters, time 5-120 sec.',
      fits='All 7 quizzes (one per chapter). No cap on how many kahoots you make.', method='FILE UPLOAD (.xlsx, the official Kahoot template, already filled in)',
      steps=['Sign in at create.kahoot.it and click Create, then Kahoot.', 'In the creator, click Add question, then Import spreadsheet.', 'Upload Kahoot_ChN_....xlsx from the chapter folder. Check the preview and click Add questions.',
             'Delete the empty first question if the creator added one. Type the title and description below in Settings. Click Save.'],
      note='If more than 10 students join, pair students on one device (one name per pair), or use Blooket or Gimkit for that round.', key='kahoot'),
 dict(app='Blooket', plan='Blooket Basic (free)', limit='Up to 60 players per game. Unlimited sets.', fits='All 7 sets.', method='FILE UPLOAD (.csv in the official Blooket template layout)',
      steps=['Sign in at dashboard.blooket.com and click Create.', 'Type the title and description below, then choose Spreadsheet Import.', 'Upload Blooket_ChN_....csv. Check the questions and click Save Set.', 'To play: open the set, click Host, pick a mode (Gold Quest, Café, Tower Defense), share the Game ID.'],
      note='Answer 3 and 4 are blank on true/false and scam-or-legit rows. That is normal.', key='blooket'),
 dict(app='Gimkit', plan='Gimkit Basic (free)', limit='Unlimited kits. Featured game modes are free for the whole class and rotate; Pro-only modes allow only 5 players on Basic. New teacher accounts get a 14-day Pro trial.',
      fits='All 7 kits.', method='FILE UPLOAD (.csv in the official Gimkit template layout)',
      steps=['Sign in at gimkit.com and click New Kit (or Create).', 'Type the title below, choose the language English, then choose Spreadsheet import.', 'Upload Gimkit_ChN_....csv. Check the questions and click Save.', 'To play: open the kit, click Play live, pick Classic or any free featured mode, share the code.'],
      note='Rows with only one wrong answer are true/false or scam-or-legit questions.', key='gimkit'),
 dict(app='Wayground (Quizizz)', plan='Wayground Starter (free)', limit='Store up to 20 resources in your library. Up to 100 students per session.', fits='All 7 quizzes (13 slots left for later chapters).',
      method='FILE UPLOAD (.xlsx, the official Wayground template, already filled in)',
      steps=['Sign in at wayground.com and click Create, then Assessment.', 'Click Start from scratch and scroll to the bottom to Import existing files, then Spreadsheet.', 'Upload Wayground_ChN_....xlsx. Review, then Continue.', 'Type the title below and click Publish. Assign as homework or host live.'],
      note='Ch 3 and Ch 4 include one open-ended question (voice mail greeting, folder plan) for the mini-projects.', key='wayground'),
 dict(app='Quizlet', plan='Quizlet free', limit='Unlimited flashcard sets. Quizlet Live and Match are free. Free students get limited Learn rounds and one practice Test per set.', fits='All 7 sets.',
      method='PASTE (copy all text from the .txt file)',
      steps=['Sign in at quizlet.com, click Create, then Flashcard set. Type the title below.', 'Click Import. Open Quizlet_ChN_....txt, select all, copy, and paste into the box.', 'Set Between term and definition: Tab. Set Between cards: New line. Click Import.', 'Click Create. For a class game: open the set, click Live (or Classic Live), choose Teams.'],
      note='Each line is one card: term, a tab, then the meaning.', key='quizlet'),
 dict(app='Wordwall', plan='Wordwall Basic (free)', limit='Create up to 3 activities, using 12 standard templates (Match up and Group sort are included). Gameshow quiz, True or false, Speed sorting and Rank order are paid.',
      fits='Only 3 activities fit. Use the 3 below. Tip: one activity can switch templates (for example Match up to Find the match) without using a new slot.', method='TYPE OR PASTE line by line (no file import)',
      steps=['Sign in at wordwall.net and click Create Activity.', 'Pick the template named in the file (Group sort or Match up).', 'Type the title and instruction. Paste each item into its own box from the .txt file.', 'Click Done. Click Share, then Publish and copy the link for Zoom chat.'],
      note='Do not make test activities: free slots are limited to 3.', key='wordwall'),
 dict(app='Padlet', plan='Padlet Neon (free)', limit='3 padlets (boards) in total. Uploads up to 20 MB each. Unlimited posts and students.', fits='Only 3 boards fit. Use the 3 below.',
      method='TYPE OR PASTE (board setup sheet; no file import)',
      steps=['Sign in at padlet.com and click Make a padlet.', 'Pick the format in the file (Wall or Columns). Paste the title and description.', 'Add the column names if listed. Pin the first post with the prompt.', 'Click Share, copy the link, and paste it in Zoom chat. Students do not need an account.'],
      note='When the course moves on, you can clear posts and reuse a board instead of making a new one.', key='padlet'),
 dict(app='Mentimeter', plan='Mentimeter Free', limit='50 participants per month in total across all presentations.', fits='2 short polls. Run #1 in October and #2 in November so each month has its own 50.',
      method='TYPE OR PASTE (slide list; no file import)',
      steps=['Sign in at mentimeter.com and click New presentation.', 'Add each slide with the type in the file (Multiple choice, Word cloud, Open ended, Scales).', 'Paste the question and options.', 'Click Present. Students go to menti.com and type the code.'],
      note='A big class can use the 50 in one session. Keep it to these 2 polls.', key='menti'),
]

def table(doc, rows, ws, head=True):
    t = doc.add_table(rows=len(rows), cols=len(rows[0])); borders(t); widths(t, ws)
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            c = t.cell(i, j); c.text = ''; p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(1)
            run = p.add_run(str(v)); run.font.size = Pt(9)
            if head and i == 0: run.bold = True; run.font.color.rgb = RGBColor(0xff, 0xff, 0xff); shade(c, '1D3557')
    para(doc, '', after=4); return t

def build(files):
    doc = Document(); sec = doc.sections[0]; sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for s in ('left_margin', 'right_margin'): setattr(sec, s, Inches(0.7))
    sec.top_margin = sec.bottom_margin = Inches(0.7)
    st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Jmal C's Digital Literacy Class · Free App Versions, Chapters 1-7 · Page "); r.font.size = Pt(8); field(fp, 'PAGE')
    para(doc, "Jmal C's Digital Literacy Class", 14, True, TEAL, 24)
    para(doc, 'Free App Versions of the Ch 1-7 Activities', 26, True, NAVY, 6)
    para(doc, 'Ready-to-import files for Kahoot, Blooket, Gimkit, Wayground, Quizlet, Wordwall, Padlet, and Mentimeter', 13, False, NAVY, 14)
    para(doc, 'Every question and card is original, written for this class from the class games and summaries. Nothing is copied from the textbook.', 10.5, after=6)
    para(doc, CHECKED, 10, italic=True, after=10)
    heading(doc, 'What Jmal does (about 2 hours once)', 2)
    for t in ['Make a free teacher account in each app (use the same email for all). No account is needed for students in any of these apps.',
              'Import the files: 7 Kahoots, 7 Blookets, 7 Gimkit kits, 7 Wayground quizzes, 7 Quizlet sets, 3 Wordwalls, 3 Padlets, 2 Mentimeters. Steps for each app are below.',
              'Type the title and description from the table for each app. Check one or two questions in the preview.',
              'Before class, open the game you need and press Host or Present. Paste the join link or code into Zoom chat.']:
        doc.add_paragraph(t, style='List Number')
    heading(doc, 'Free plan limits that decide what fits', 2)
    table(doc, [['App', 'Free plan', 'Limit', 'What fits']] + [[a['app'], a['plan'], a['limit'], a['fits']] for a in APPS], [Inches(1.0), Inches(1.2), Inches(2.9), Inches(2.0)])
    heading(doc, 'Which app for each activity', 2)
    rows = [['Ch', 'Activity', 'App', 'Use this', 'How to run it']]
    for (n, i), (app, what, how) in sorted(MAP.items()): rows.append([n, A[n][i - 1]['t'], app, what, how])
    table(doc, rows, [Inches(0.35), Inches(1.8), Inches(0.9), Inches(2.2), Inches(1.85)])
    for a in APPS:
        doc.add_page_break(); heading(doc, a['app'], 1)
        label(doc, 'Free plan:', f"{a['plan']}. {a['limit']}"); label(doc, 'What fits:', a['fits']); label(doc, 'Import method:', a['method'])
        heading(doc, 'Steps', 3)
        for s in a['steps']: doc.add_paragraph(s, style='List Number')
        label(doc, 'Note:', a['note'])
        heading(doc, 'Files, titles, and descriptions', 3)
        k = a['key']
        if k in ('kahoot', 'blooket', 'gimkit', 'wayground', 'quizlet'):
            rows = [['Ch', 'File (in Import_Files)', 'Items', 'Title', 'Description']]
            for n in range(1, 8):
                p, items = files[n][k]; rows.append([n, REL(p).replace('Import_Files/', ''), len(items), M.title(n), M.desc(n)])
            table(doc, rows, [Inches(0.35), Inches(2.4), Inches(0.5), Inches(1.6), Inches(2.25)])
        else:
            src = {'wordwall': WORDWALL, 'padlet': PADLET, 'menti': MENTI}[k]
            rows = [['#', 'Ch', 'File (in Import_Files)', 'Title', 'Used for']]
            for x in src:
                p = [q for q in files[x['ch']][k] if f"_{x['n']}_" in os.path.basename(q)][0]
                use = [A[c][i - 1]['t'] for (c, i), v in MAP.items() if v[0].lower().startswith(a['app'].lower()[:5]) and c == x['ch']]
                rows.append([x['n'], x['ch'], REL(p).replace('Import_Files/', ''), x['title'], ', '.join(use)])
            table(doc, rows, [Inches(0.3), Inches(0.35), Inches(2.6), Inches(1.9), Inches(1.95)])
            for x in src:
                p = [q for q in files[x['ch']][k] if f"_{x['n']}_" in os.path.basename(q)][0]
                heading(doc, f"{a['app']} #{x['n']}: {x['title']}", 3)
                for line in open(p, encoding='utf-8').read().splitlines()[1:]:
                    para(doc, line, 9.5, after=0)
    doc.add_page_break(); heading(doc, 'Folder map', 1)
    for d in sorted(glob.glob(os.path.join(OUT, 'Import_Files', '*'))):
        para(doc, os.path.basename(d) + '/', 10.5, True, NAVY, 1)
        for f in sorted(os.listdir(d)): para(doc, '    ' + f, 9.5, after=0)
    para(doc, '', after=6)
    para(doc, 'Blank official templates (Kahoot, Wayground, Gimkit) are in _Official_Blank_Templates, so you can add your own questions in the same format later.', 10, italic=True)
    doc.save(os.path.join(OUT, NAME + '.docx'))
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', OUT, os.path.join(OUT, NAME + '.docx')], check=True, capture_output=True)
    z = os.path.join(OUT, 'Free_App_Import_Files_Ch1-7.zip')
    with zipfile.ZipFile(z, 'w', zipfile.ZIP_DEFLATED) as zf:
        for f in sorted(glob.glob(os.path.join(OUT, 'Import_Files', '**', '*'), recursive=True)):
            if os.path.isfile(f): zf.write(f, os.path.relpath(f, OUT))
        zf.write(os.path.join(OUT, NAME + '.pdf'), NAME + '.pdf')
    return z

if __name__ == '__main__':
    z = build(M.build()); print(z, os.path.getsize(z))
