"""Build Fun_Activities_Ch1-7.docx (python-docx) and .pdf (LibreOffice) from tools/activities.py."""
import os, sys, subprocess, shutil
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from activities import A, CARDS
from chapters import CHAPTERS
OUT = '/workspace/BCC_IC3_Class/Fun_Activities_Ch1-7'
NAVY, TEAL, GOLD = RGBColor(0x1d, 0x35, 0x57), RGBColor(0x0f, 0x7c, 0x83), RGBColor(0xb8, 0x86, 0x00)

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement('w:shd'); s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), hexcolor); tcPr.append(s)

def borders(table, style='single', color='9AA8B8', sz=8):
    tbl = table._tbl; tblPr = tbl.tblPr; b = OxmlElement('w:tblBorders')
    for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        x = OxmlElement(f'w:{e}'); x.set(qn('w:val'), style); x.set(qn('w:sz'), str(sz)); x.set(qn('w:color'), color); b.append(x)
    tblPr.append(b)

def widths(table, ws):
    table.autofit = False
    for i, w in enumerate(ws):
        table.columns[i].width = w
        for c in table.columns[i].cells: c.width = w

def field(par, instr):
    r = par.add_run(); f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin'); r._r.append(f1)
    r2 = par.add_run(); it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr; r2._r.append(it)
    r3 = par.add_run(); f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'end'); r3._r.append(f2)

def para(doc, text='', size=11, bold=False, color=None, after=4, italic=False, align=None):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    if align: p.alignment = align
    if text:
        r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.italic = italic
        if color: r.font.color.rgb = color
    return p

def label(doc, lab, text):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(lab + ' '); r.bold = True; r.font.color.rgb = TEAL; p.add_run(text)
    return p

def heading(doc, text, lvl):
    h = doc.add_heading(text, lvl)
    for r in h.runs: r.font.color.rgb = NAVY
    return h

def build():
    doc = Document()
    sec = doc.sections[0]; sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for s in ('left_margin', 'right_margin'): setattr(sec, s, Inches(0.8))
    sec.top_margin, sec.bottom_margin = Inches(0.7), Inches(0.7)
    st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(11); st.paragraph_format.space_after = Pt(2); st.paragraph_format.line_spacing = 1.1
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    for n, sz in (('Heading 1', 20), ('Heading 2', 14), ('Heading 3', 12)):
        doc.styles[n].font.name = 'Calibri'; doc.styles[n].font.size = Pt(sz)
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Jmal C's Digital Literacy Class · Fun Activities, Chapters 1-7 · Page "); r.font.size = Pt(8); r.font.color.rgb = RGBColor(0x7a, 0x87, 0x96)
    field(fp, 'PAGE')
    # Cover
    para(doc, "Jmal C's Digital Literacy Class", 14, True, TEAL, 30)
    para(doc, 'Fun Activities for Chapters 1-7', 30, True, NAVY, 6)
    para(doc, 'Teacher Activity Pack', 18, False, NAVY, 18)
    para(doc, 'IC3 GS5 Computing Fundamentals · Fall 2026 · Mondays and Wednesdays, 6:00 to 9:00 PM ET', 11, False, None, 4)
    para(doc, 'Week 1 is in person. Weeks 2 to 6 are on Zoom. Every activity has an in-person version and a Zoom version.', 11, False, None, 18)
    para(doc, 'How to use this pack', 13, True, NAVY, 4)
    for t in ["Each chapter has 3 or 4 activities of 10 to 25 minutes. Pick 1 or 2 per class to break up the 3 hours.",
              "Every activity lists the objective, time, materials, in-person and Zoom versions, steps, a debrief with answers, and an ESL tip.",
              "Printable cards and handouts come right after each chapter's activities. Print one set per team, or show them on screen for Zoom.",
              "Students can practice the same ideas in the class site games: open a chapter on the student home page and tap Games.",
              "All activities are original for this class. They match the chapter topics but do not copy the textbook."]:
        p = doc.add_paragraph(t, style='List Bullet'); p.paragraph_format.space_after = Pt(2)
    para(doc, '', after=6)
    para(doc, 'Zoom tips for every activity', 13, True, NAVY, 4)
    for t in ["Breakout rooms of 3 work best. Name one person the screen sharer and one the reporter.",
              "Use the chat waterfall: everyone types an answer, then presses Enter together on \"3, 2, 1, send\".",
              "Use reactions as votes (thumbs-up, heart) so quieter students can answer without speaking.",
              "Put instructions in chat before opening breakout rooms. Broadcast a 1-minute warning."]:
        p = doc.add_paragraph(t, style='List Bullet'); p.paragraph_format.space_after = Pt(2)
    para(doc, 'ESL tips for every activity', 13, True, NAVY, 4)
    for t in ["Model first, then do it together, then let teams do it alone.", "Write key words with a picture on the board or slide.",
              "Give sentence frames. Accept one-word answers at first.", "Mix language levels in teams and give each person a job."]:
        p = doc.add_paragraph(t, style='List Bullet'); p.paragraph_format.space_after = Pt(2)
    doc.add_page_break()
    # Overview table
    heading(doc, 'All activities at a glance', 1)
    t = doc.add_table(rows=1, cols=4); borders(t); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(['Chapter', 'Activity', 'Type', 'Minutes']):
        c = t.rows[0].cells[i]; c.text = ''; rr = c.paragraphs[0].add_run(h); rr.bold = True; rr.font.color.rgb = RGBColor(255, 255, 255); shade(c, '1D3557')
    for c in CHAPTERS:
        for a in A[c['n']]:
            row = t.add_row().cells; row[0].text = f"{c['n']}. {c['title']}"; row[1].text = a['t']; row[2].text = a['kind']; row[3].text = str(a['min'])
    widths(t, (Inches(1.9), Inches(2.6), Inches(1.6), Inches(0.8)))
    # Chapters
    for c in CHAPTERS:
        n = c['n']; doc.add_page_break()
        heading(doc, f"Chapter {n}: {c['title']}", 1)
        para(doc, 'Activities: ' + ' · '.join(a['t'] for a in A[n]), 10, False, TEAL, 8, italic=True)
        for k, a in enumerate(A[n], 1):
            if k > 1: para(doc, '', after=2)
            heading(doc, f"{n}.{k} {a['t']}", 2)
            info = doc.add_table(rows=2, cols=3); borders(info, color='C9D4E1')
            for i, (h, v) in enumerate((('Type', a['kind']), ('Time', f"{a['min']} minutes"), ('Objective', a['obj']))):
                hc = info.rows[0].cells[i]; hc.text = ''; rr = hc.paragraphs[0].add_run(h); rr.bold = True; rr.font.size = Pt(9); rr.font.color.rgb = NAVY; shade(hc, 'EEF3F8')
                vc = info.rows[1].cells[i]; vc.text = ''; r2 = vc.paragraphs[0].add_run(v); r2.font.size = Pt(10)
            widths(info, (Inches(1.4), Inches(1.0), Inches(4.5)))
            para(doc, '', after=0)
            label(doc, 'Materials:', a['mat'])
            label(doc, 'In person:', a['inp'])
            label(doc, 'On Zoom:', a['zoom'])
            p = doc.add_paragraph(); r = p.add_run('Steps'); r.bold = True; r.font.color.rgb = TEAL; p.paragraph_format.space_after = Pt(1)
            for si, s in enumerate(a['steps'], 1):
                q = doc.add_paragraph(); q.paragraph_format.space_after = Pt(1); q.paragraph_format.left_indent = Inches(0.3); q.paragraph_format.first_line_indent = Inches(-0.22)
                rn = q.add_run(f'{si}.  '); rn.bold = True; q.add_run(s)
            label(doc, 'Debrief and answer key:', a['key'])
            box = doc.add_table(rows=1, cols=1); borders(box, color='FFCC33', sz=12); cc = box.rows[0].cells[0]; shade(cc, 'FFF8E1'); cc.text = ''
            rr = cc.paragraphs[0].add_run('ESL tip: '); rr.bold = True; cc.paragraphs[0].add_run(a['esl'])
            if a.get('cards'):
                para(doc, f"Printable: {CARDS[a['cards']]['t']} (see the card pages after this chapter).", 9, False, TEAL, 2, italic=True)
        # Cards for this chapter
        for a in A[n]:
            cid = a.get('cards')
            if not cid: continue
            cs = CARDS[cid]; doc.add_page_break()
            heading(doc, f"Printable · Chapter {n} · {cs['t']}", 2)
            para(doc, f"For activity {n}.{A[n].index(a) + 1}, {a['t']}. " + ('Cut along the lines.' if cs['cols'] > 1 else 'Print one per student or team.'), 9, False, TEAL, 6, italic=True)
            if cs.get('bingo'):
                tb = doc.add_table(rows=2, cols=2); borders(tb, 'dashed', '7A8796')
                for bi, board in enumerate(cs['bingo']):
                    cell = tb.rows[bi // 2].cells[bi % 2]; cell.text = ''
                    pp = cell.paragraphs[0]; rr = pp.add_run(f'EXTENSION BINGO · Board {bi + 1}'); rr.bold = True; rr.font.color.rgb = NAVY
                    g = cell.add_table(rows=3, cols=3); borders(g, color='1D3557', sz=10)
                    for j, v in enumerate(board):
                        gc = g.rows[j // 3].cells[j % 3]; gc.text = ''; gp = gc.paragraphs[0]; gp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        gr = gp.add_run(v); gr.bold = True; gr.font.size = Pt(18); gc.height = Cm(2)
                        if v == 'FREE': shade(gc, 'FFF3C4')
                para(doc, 'Word bank: .docx .xlsx .pptx .pdf .jpg .png .mp3 .mp4 .exe .zip .txt .wav', 10, after=2)
                continue
            cols = cs['cols']; items = cs['items']; rows = (len(items) + cols - 1) // cols
            tb = doc.add_table(rows=rows, cols=cols); borders(tb, 'dashed' if cols > 1 else 'single', '7A8796' if cols > 1 else 'C9D4E1')
            for j, v in enumerate(items):
                cell = tb.rows[j // cols].cells[j % cols]; cell.text = ''; pp = cell.paragraphs[0]
                if cols > 1: pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                rr = pp.add_run(v); rr.font.size = Pt(16 if cols == 3 and len(v) < 30 else 13 if cols > 1 else 12)
                if cs.get('mono'): rr.font.name = 'Courier New'
                if cols > 1: rr.bold = len(v) < 40
                pp.paragraph_format.space_before = Pt(14 if cols > 1 else 6); pp.paragraph_format.space_after = Pt(14 if cols > 1 else 6)
    doc.save(OUT + '.docx')
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', os.path.dirname(OUT), OUT + '.docx'], check=True, capture_output=True, timeout=180)
    pages = subprocess.run(['pdfinfo', OUT + '.pdf'], capture_output=True, text=True).stdout
    print(OUT, [l for l in pages.splitlines() if l.startswith('Pages')])

def chapter_pages():
    """First PDF page of each chapter's activities (for teacher links)."""
    import re
    out = {}
    n_pages = int(re.search(r'Pages:\s+(\d+)', subprocess.run(['pdfinfo', OUT + '.pdf'], capture_output=True, text=True).stdout).group(1))
    for pg in range(1, n_pages + 1):
        t = subprocess.run(['pdftotext', '-f', str(pg), '-l', str(pg), '-layout', OUT + '.pdf', '-'], capture_output=True, text=True).stdout
        for c in CHAPTERS:
            if c['n'] not in out and re.search(rf"^\s*Chapter {c['n']}: {re.escape(c['title'])}\s*$", t, re.M): out[c['n']] = pg
    return out

if __name__ == '__main__':
    build(); print(chapter_pages())
